import os
import sqlite3
import tempfile
from pathlib import Path

from docx import Document

from core.common_data_area import CommonDataArea
from core.executor import Executor
from core.db_schema import init_db
from tools.tool_registry import call_tool, get_tool, sync_tools_to_db


def test_export_file_writes_csv_under_storage_subfolder():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)

        result = call_tool(
            'export_file',
            {
                'output_filename': 'report',
                'format': 'csv',
                'subfolder': 'exports/reports',
                'columns': [
                    {'key': 'name', 'header': 'Name'},
                    {'key': 'amount', 'header': 'Amount'},
                ],
                'rows': [
                    {'name': 'Alpha', 'amount': 10},
                    {'name': 'Beta', 'amount': 20},
                ],
            },
        )

        assert result['success'] is True
        file_path = Path(result['file_path'])
        assert file_path.exists()
        assert file_path.suffix.lower() == '.csv'
        assert file_path.parent == Path(storage_dir, 'exports', 'reports')
        content = file_path.read_text(encoding='utf-8')
        assert 'Name,Amount' in content
        assert 'Alpha,10' in content


def test_render_image_writes_png_under_storage_subfolder():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)

        result = call_tool(
            'render_image',
            {
                'output_filename': 'summary_card',
                'format': 'png',
                'mode': 'canvas',
                'subfolder': 'generated/cards',
                'width': 640,
                'height': 360,
                'background': {'color': '#ffffff'},
                'data': {'title': 'Weekly Summary'},
                'elements': [
                    {'type': 'rect', 'x': 20, 'y': 20, 'width': 600, 'height': 320, 'fill': '#f4f7fb'},
                    {'type': 'text', 'x': 40, 'y': 40, 'width': 300, 'height': 50, 'text': '{{title}}', 'font_size': 24, 'bold': True},
                ],
            },
        )

        assert result['success'] is True
        file_path = Path(result['file_path'])
        assert file_path.exists()
        assert file_path.suffix.lower() == '.png'
        assert file_path.parent == Path(storage_dir, 'generated', 'cards')
        assert result['width'] == 640
        assert result['height'] == 360


def test_sync_tools_to_db_registers_export_and_render_tools():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as tmpdir:
        db_path = os.path.join(tmpdir, 'tools.db')
        cda.set_setting('sqlite_db_path', db_path)
        init_db()
        sync_tools_to_db(cda)

        conn = sqlite3.connect(db_path)
        try:
            cur = conn.cursor()
            cur.execute("SELECT name FROM ToolList WHERE name IN (?, ?) ORDER BY name", ('export_file', 'render_image'))
            names = [row[0] for row in cur.fetchall()]
        finally:
            conn.close()

        assert names == ['export_file', 'render_image']
        assert get_tool('export_file') is not None
        assert get_tool('render_image') is not None


def test_executor_tracks_created_files_for_export_tool():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)
        executor = Executor(cda)
        result = executor._execute_tool_call(
            agent_name='tester',
            loop_idx=0,
            tool_name='export_file',
            parameters={
                'output_filename': 'tracked_file',
                'format': 'csv',
                'subfolder': 'tracked',
                'columns': [{'key': 'value', 'header': 'Value'}],
                'rows': [{'value': 'ok'}],
            },
            status_cb=None,
        )

        assert result['success'] is True
        tool_data = cda.get_memory('tool_data')
        assert tool_data['success'] is True
        created_files = cda.get_memory('created_files')
        assert isinstance(created_files, list)
        assert created_files
        assert created_files[-1]['file_path'] == result['file_path']
        assert cda.get_memory('last_created_file')['file_path'] == result['file_path']


def test_toollist_sync_includes_rich_metadata_for_export_tool():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as tmpdir:
        db_path = os.path.join(tmpdir, 'tools_meta.db')
        cda.set_setting('sqlite_db_path', db_path)
        init_db()
        sync_tools_to_db(cda)

        conn = sqlite3.connect(db_path)
        try:
            cur = conn.cursor()
            cur.execute(
                "SELECT description, input_schema, output_schema, example_call FROM ToolList WHERE name=?",
                ('export_file',),
            )
            row = cur.fetchone()
        finally:
            conn.close()

        assert row is not None
        description, input_schema, output_schema, example_call = row
        assert 'configured file storage path' in description
        assert 'output_filename' in input_schema
        assert 'bytes written' in output_schema.lower()
        assert 'attendance_report' in example_call


def test_export_file_accepts_column_aliases():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)

        result = call_tool(
            'export_file',
            {
                'output_filename': 'alias_report',
                'format': 'csv',
                'subfolder': 'exports',
                'columns': [
                    {'data_key': 'name', 'header': 'Name'},
                    {'data_property': 'amount', 'header': 'Amount'},
                    {'name': 'status', 'header': 'Status'},
                ],
                'rows': [
                    {'name': 'Alpha', 'amount': 10, 'status': 'open'},
                ],
            },
        )

        assert result['success'] is True
        content = Path(result['file_path']).read_text(encoding='utf-8')
        assert 'Name,Amount,Status' in content
        assert 'Alpha,10,open' in content


def test_executor_builds_tool_examples_from_rich_parameter_metadata():
    cda = CommonDataArea()
    cda.reset()
    executor = Executor(cda)
    rendered = executor._build_detailed_tool_list(
        [
            {
                'name': 'export_file',
                'description': 'Create a document file.',
                'input_schema': '{"parameters": [{"name": "output_filename", "type": "str", "required": true, "description": "Target file name."}, {"name": "format", "type": "str", "required": true, "description": "Output format."}]}',
                'example_call': '{"action": {"type": "tool_call", "tool_request": {"tool_name": "export_file", "parameters": {"output_filename": "report.pdf", "format": "pdf"}}}}',
            }
        ]
    )

    assert 'output_filename (str, required): Target file name.' in rendered
    assert 'format (str, required): Output format.' in rendered
    assert '"output_filename": "report.pdf"' in rendered
    assert "{'name': 'output_filename'" not in rendered

def test_export_file_uses_task_report_layout_for_task_pdf():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)

        result = call_tool(
            'export_file',
            {
                'output_filename': 'tasks_for_bijumon',
                'format': 'pdf',
                'subfolder': 'task_reports',
                'title': 'Tasks for Bijumon',
                'columns': [
                    {'key': 'id', 'header': 'ID'},
                    {'key': 'title', 'header': 'Title'},
                    {'key': 'description', 'header': 'Description'},
                    {'key': 'status', 'header': 'Status'},
                    {'key': 'scheduled_date', 'header': 'Scheduled Date'},
                    {'key': 'due_date', 'header': 'Due Date'},
                    {'key': 'priority', 'header': 'Priority'},
                    {'key': 'assigned_to', 'header': 'Assigned To'},
                ],
                'rows': [
                    {
                        'id': 1,
                        'title': 'Refactor NEERS app',
                        'description': 'Verify filters, timings, electrotonus, and remove logo.',
                        'status': 'open',
                        'scheduled_date': '2026-03-02',
                        'due_date': '2026-03-15 23:59:59',
                        'priority': 'medium',
                        'assigned_to': 'Bijumon',
                    }
                ],
            },
        )

        assert result['success'] is True
        assert Path(result['file_path']).exists()
        assert result['meta']['layout'] == 'task_report'
        assert result['meta']['page_orientation'] == 'landscape'

def test_export_file_builds_docx_from_sections_table_payload():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)

        result = call_tool(
            'export_file',
            {
                'output_filename': 'section_tasks.docx',
                'format': 'docx',
                'subfolder': 'tasks',
                'sections': [
                    {'type': 'heading', 'level': 1, 'content': 'Tasks Assigned to Bijumon'},
                    {
                        'type': 'table',
                        'columns': [
                            {'key': 'id', 'header': 'ID'},
                            {'key': 'title', 'header': 'Title'},
                            {'key': 'description', 'header': 'Description'},
                            {'key': 'status', 'header': 'Status'},
                            {'key': 'scheduled_date', 'header': 'Scheduled Date'},
                            {'key': 'due_date', 'header': 'Due Date'},
                            {'key': 'priority', 'header': 'Priority'},
                            {'key': 'assigned_to', 'header': 'Assigned To'},
                        ],
                        'rows': [
                            {
                                'id': 1,
                                'title': 'Refactor NEERS app',
                                'description': 'Verify filters, timings, electrotonus, and remove logo.',
                                'status': 'open',
                                'scheduled_date': '2026-03-02',
                                'due_date': '2026-03-15 23:59:59',
                                'priority': 'medium',
                                'assigned_to': 'Bijumon',
                            },
                            {
                                'id': 2,
                                'title': 'Create export tool',
                                'description': 'Implement file export and image rendering tools.',
                                'status': 'closed',
                                'scheduled_date': '2026-03-08',
                                'due_date': None,
                                'priority': 'high',
                                'assigned_to': 'Bijumon',
                            },
                        ],
                    },
                ],
            },
        )

        assert result['success'] is True
        file_path = Path(result['file_path'])
        assert file_path.exists()
        assert result['meta']['row_count'] == 2
        assert result['meta']['column_count'] == 8
        assert result['meta']['layout'] == 'task_report'
        assert result['meta']['page_orientation'] == 'landscape'

        doc = Document(str(file_path))
        assert doc.tables
        assert doc.tables[0].style.name == 'Table Grid'
        assert doc.tables[0].rows[0].cells[0].text == 'ID'
        assert doc.tables[0].rows[1].cells[1].text == 'Refactor NEERS app'


def test_export_file_uses_task_report_layout_from_sections_table_pdf():
    cda = CommonDataArea()
    cda.reset()

    with tempfile.TemporaryDirectory() as storage_dir:
        cda.set_setting('file_storage_path', storage_dir)

        result = call_tool(
            'export_file',
            {
                'output_filename': 'section_tasks.pdf',
                'format': 'pdf',
                'subfolder': 'tasks',
                'sections': [
                    {'type': 'heading', 'level': 1, 'content': 'Tasks Assigned to Bijumon'},
                    {
                        'type': 'table',
                        'columns': [
                            {'key': 'id', 'header': 'ID'},
                            {'key': 'title', 'header': 'Title'},
                            {'key': 'description', 'header': 'Description'},
                            {'key': 'status', 'header': 'Status'},
                            {'key': 'scheduled_date', 'header': 'Scheduled Date'},
                            {'key': 'due_date', 'header': 'Due Date'},
                            {'key': 'priority', 'header': 'Priority'},
                            {'key': 'assigned_to', 'header': 'Assigned To'},
                        ],
                        'rows': [
                            {
                                'id': 1,
                                'title': 'Refactor NEERS app',
                                'description': 'Verify filters, timings, electrotonus, and remove logo.',
                                'status': 'open',
                                'scheduled_date': '2026-03-02',
                                'due_date': '2026-03-15 23:59:59',
                                'priority': 'medium',
                                'assigned_to': 'Bijumon',
                            }
                        ],
                    },
                ],
            },
        )

        assert result['success'] is True
        assert Path(result['file_path']).exists()
        assert result['meta']['layout'] == 'task_report'
        assert result['meta']['page_orientation'] == 'landscape'
        assert result['meta']['row_count'] == 1
        assert result['meta']['column_count'] == 8


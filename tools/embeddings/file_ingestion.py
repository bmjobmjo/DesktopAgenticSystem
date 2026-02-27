
import os
import shutil
import json
import hashlib
import sqlite3
from pathlib import Path
from typing import List, Dict, Any, Optional, Callable

from core.common_data_area import CommonDataArea
from core.db_schema import DB_PATH
from execution_logger import log_execution_step, log_exception

# -----------------------------------------------------------------------------
# Lazy Imports for File Processing & Embeddings
# -----------------------------------------------------------------------------
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pptx
except ImportError:
    pptx = None

try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None

try:
    from sentence_transformers import SentenceTransformer
    _EMBEDDING_MODEL = None
except ImportError:
    _EMBEDDING_MODEL = None
    SentenceTransformer = None

def _load_sentence_transformer(model_id_or_path: str, local_files_only: bool = False):
    """
    Load sentence-transformer model with compatibility fallback for older package versions.
    """
    try:
        return SentenceTransformer(model_id_or_path, local_files_only=local_files_only)
    except TypeError:
        # Older sentence-transformers may not support local_files_only argument.
        return SentenceTransformer(model_id_or_path)

def get_embedding_model():
    global _EMBEDDING_MODEL
    if _EMBEDDING_MODEL is None:
        if SentenceTransformer is None:
            log_execution_step('EMBEDDING_INIT', "SentenceTransformer class is not available (Import failed).")
            return None

        cda = CommonDataArea()
        model_name = str(cda.get_setting('embedding_model_name', 'all-MiniLM-L6-v2') or 'all-MiniLM-L6-v2').strip()
        model_path = str(cda.get_setting('embedding_model_path', '') or '').strip()
        local_only = bool(cda.get_setting('embedding_local_files_only', False))

        last_error = None
        candidates = []
        if model_path:
            candidates.append((model_path, True))
        candidates.append((model_name, local_only))
        # Final fallback: cached local copy by model name, no network.
        candidates.append((model_name, True))

        seen = set()
        for model_id_or_path, local_files_only in candidates:
            key = (model_id_or_path, bool(local_files_only))
            if key in seen:
                continue
            seen.add(key)
            try:
                mode = "local-only" if local_files_only else "online/cached"
                log_execution_step('EMBEDDING_INIT', f"Loading SentenceTransformer model '{model_id_or_path}' ({mode})...")
                _EMBEDDING_MODEL = _load_sentence_transformer(model_id_or_path, local_files_only=local_files_only)
                log_execution_step('EMBEDDING_INIT', f"Model loaded successfully from '{model_id_or_path}'.")
                break
            except Exception as e:
                last_error = e
                log_exception(
                    'EMBEDDING_LOAD_ATTEMPT_FAILED',
                    e,
                    {'model': model_id_or_path, 'local_files_only': local_files_only}
                )

        if _EMBEDDING_MODEL is None:
            print(
                "Failed to load embedding model. Configure one of:\n"
                "1) settings.user_config.json -> embedding_model_path: '<local_model_folder>'\n"
                "2) settings.user_config.json -> embedding_model_name: 'all-MiniLM-L6-v2' (internet/cache required)\n"
                "3) settings.user_config.json -> embedding_local_files_only: true (use local cache only)\n"
                f"Last error: {last_error}"
            )
    return _EMBEDDING_MODEL

# -----------------------------------------------------------------------------
# Optimization: Text Extraction Logic
# -----------------------------------------------------------------------------
def _extract_text(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    
    if ext == '.pdf':
        if not PyPDF2: return "[PyPDF2 not installed]"
        text = ""
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"
        except Exception as e:
            text = f"[Error reading PDF: {e}]"
        return text
        
    elif ext in ['.docx', '.doc']:
        if not docx: return "[python-docx not installed]"
        try:
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            return f"[Error reading DOCX: {e}]"
            
    elif ext in ['.xlsx', '.xls']:
        if not openpyxl: return "[openpyxl not installed]"
        text = ""
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                text += f"Sheet: {sheet}\n"
                for row in ws.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    text += row_text + "\n"
        except Exception as e:
            text = f"[Error reading Excel: {e}]"
        return text

    elif ext in ['.pptx', '.ppt']:
        if not pptx: return "[python-pptx not installed]"
        text = ""
        try:
            prs = pptx.Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        except Exception as e:
            text = f"[Error reading PPT: {e}]"
        return text

    elif ext in ['.png', '.jpg', '.jpeg', '.bmp', '.tiff']:
        if not pytesseract or not Image: return "[pytesseract/Pillow not installed]"
        try:
            return pytesseract.image_to_string(Image.open(file_path))
        except Exception as e:
            return f"[OCR Error: {e}]"

    else:
        # Default to plain text
        try:
            return file_path.read_text(encoding='utf-8', errors='replace')
        except Exception as e:
            return f"[Error reading text: {e}]"

def _compute_hash(file_path: Path) -> str:
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for b in iter(lambda: f.read(4096), b""):
            sha256.update(b)
    return sha256.hexdigest()

def _chunk_text(text: str, chunk_size: int = 500) -> List[str]:
    return [text[i:i+chunk_size] for i in range(0, len(text), chunk_size)]

def _get_embeddings_db_path() -> Path:
    cda = CommonDataArea()
    configured = str(cda.get_setting('sqlite_db_path', '') or '').strip()
    if configured:
        return Path(configured).resolve()
    return Path(DB_PATH).resolve()

# -----------------------------------------------------------------------------
# Main Tool Functions
# -----------------------------------------------------------------------------

def ingest_file(file_path: str, category: str = None, user_description: str = None, status_callback: Optional[Callable[[str], None]] = None) -> Dict[str, Any]:
    """
    Ingest a file into the knowledge base (copy, index, and embed).

    Process:
    1. Copies the file to the configured 'file_storage_path'.
    2. Extracts text content (PDF, DOCX, XLSX, PPTX, Images via OCR).
    3. Generates a description using LLM (incorporating user_description if provided).
    4. Creates vector embeddings for the description and content chunks.
    5. Stores metadata in the Database.

    Args:
        file_path (str): Absolute path to the source file.
        category (str, optional): Metadata category tag.
        user_description (str, optional): A user-provided description to guide the LLM summary.
        status_callback (callable, optional): UI status update handler.

    Returns:
        Dict: Success status and metadata including 'file_id' and 'description'.
    """
    if status_callback:
        status_callback(f"Processing {os.path.basename(file_path)}...")
        
    cda = CommonDataArea()
    source_path = Path(file_path).resolve()
    
    if not source_path.exists():
        return {"success": False, "error": f"Source file not found: {file_path}"}

    # 1. Read Settings
    storage_setting = cda.get_setting('file_storage_path', 'storage/files')
    # Resolve storage path relative to project root (CWD in most cases here)
    storage_dir = Path(storage_setting).resolve()
    storage_dir.mkdir(parents=True, exist_ok=True)

    # 2. Copy File
    dest_path = storage_dir / source_path.name
    # Handle naming collision if needed, but for now overwrite or simple copy
    if dest_path.exists():
        # Checking hash to see if identical?
        if _compute_hash(source_path) == _compute_hash(dest_path):
            pass # Same file
        else:
            # Append timestamp or incremental
            stem = source_path.stem
            ext = source_path.suffix
            import time
            dest_path = storage_dir / f"{stem}_{int(time.time())}{ext}"
    
    shutil.copy2(source_path, dest_path)
    final_path = dest_path

    # 3. Create File Entry in DB
    file_hash = _compute_hash(final_path)
    file_size = final_path.stat().st_size
    
    # Store relative path in DB for portability
    try:
        db_path_str = os.path.relpath(final_path, os.getcwd())
    except Exception:
        db_path_str = str(final_path)
    
    log_execution_step('INGEST', f"Final storage path (relative): {db_path_str}")

    db_path = _get_embeddings_db_path()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(str(db_path), timeout=20) as conn:
        cursor = conn.cursor()
        
        # Check if this exact file (by path) exists?
        cursor.execute("SELECT id FROM Files WHERE file_path = ?", (db_path_str,))
        existing = cursor.fetchone()
        
        if existing:
            file_id = existing[0]
            # Update metadata
            cursor.execute("UPDATE Files SET file_hash=?, file_size=?, upload_date=CURRENT_TIMESTAMP WHERE id=?", 
                           (file_hash, file_size, file_id))
            # Clear old embeddings/expenses
            cursor.execute("DELETE FROM FileEmbeddings WHERE file_id=?", (file_id,))
            cursor.execute("DELETE FROM Expenses WHERE file_id=?", (file_id,))
        else:
            cursor.execute("INSERT INTO Files (file_path, filename, file_type, file_size, file_hash) VALUES (?, ?, ?, ?, ?)",
                           (db_path_str, final_path.name, final_path.suffix, file_size, file_hash))
            file_id = cursor.lastrowid
            
        conn.commit()

        # Phase 1: Text Extraction
        log_execution_step('INGEST', f"Starting extraction for {final_path.name}")
        if status_callback:
            status_callback(f"Extracting text from {final_path.name}...")
        
        try:
            text_content = _extract_text(final_path)
            log_execution_step('INGEST', f"Extracted {len(text_content)} characters.")
        except Exception as e:
            log_exception('INGEST_ERROR', e, f"Text extraction failed for {final_path.name}")
            raise
            
        # Phase 2: Embed Document Content first
        model = get_embedding_model()
        if not model:
            err = "Embedding model (SentenceTransformers) could not be loaded."
            log_execution_step('INGEST_ERROR', err)
            raise RuntimeError(err)

        chunks = _chunk_text(text_content, 1000)
        log_execution_step('INGEST', f"Document split into {len(chunks)} chunks. Starting vectorization.")
        
        try:
            for i, chk in enumerate(chunks):
                if status_callback:
                    status_callback(f"Embedding chunk {i+1}/{len(chunks)}...")
                
                vec = model.encode([chk])[0]
                cursor.execute(
                    "INSERT INTO FileEmbeddings (file_id, chunk_index, chunk_type, content_chunk, embedding_vector) VALUES (?, ?, ?, ?, ?)",
                    (file_id, i, 'text', chk, json.dumps(vec.tolist()))
                )
                if i == 0 or i == len(chunks) - 1 or (i+1) % 5 == 0:
                    log_execution_step('INGEST_DB', f"Saved chunk {i+1}/{len(chunks)} to FileEmbeddings.")

            # Commit content embeddings immediately so they are saved even if LLM fails later
            conn.commit()
            log_execution_step('INGEST', "Primary content embeddings committed to database.")
            
        except Exception as e:
            log_exception('INGEST_DB_ERROR', e, "Failed to save content embeddings.")
            raise

        # Phase 3: LLM Enrichment (Description & Expense Category)
        description = "No description generated."
        is_expense = False
        expense_data = {}
        
        llm = cda.get_runtime('llm_client')
        if llm:
            snippet = text_content[:2000]
            log_execution_step('INGEST', "Requesting LLM analysis for description & classification...")
            user_context = ""
            if user_description:
                user_context = f"\n        User provided description/context: '{user_description}'\n        Please incorporate this user context into the final description."

            prompt = f"""
            Analyze the following file content and provide:
            1. A brief 1-sentence description.{user_context}
            2. Boolean 'is_expense': true if this looks like a receipt/invoice/bill that tracks spending.
            3. If is_expense, extract JSON: {{ "amount": float, "currency": "USD", "vendor": "str", "date": "YYYY-MM-DD", "category": "str" }}
            
            Content:
            {snippet}
            
            Return JSON only.
            """
            try:
                if status_callback:
                    status_callback("Analyzing metadata with LLM...")
                response = llm.generate(prompt, agent_name="Ingestor", user_prompt="Ingest File")
                import re
                json_match = re.search(r'\{.*\}', response, re.DOTALL)
                if json_match:
                    data = json.loads(json_match.group(0))
                    description = data.get('description', description)
                    is_expense = data.get('is_expense', False)
                    if is_expense:
                        expense_data = data
                log_execution_step('INGEST', f"LLM metadata analysis successful: {description[:50]}...")
            except Exception as e:
                log_exception('INGEST_LLM_ERROR', e, "LLM Enrichment failed - continuing without description.")

        # Phase 4: Final Updates (Description & Description Embedding)
        try:
            # Update File record
            cursor.execute("UPDATE Files SET description = ? WHERE id = ?", (description, file_id))
            log_execution_step('INGEST_DB', f"Updated file {file_id} description.")

            # Create special 'description' embedding for boosted search
            desc_vec = model.encode([description])[0]
            cursor.execute(
                "INSERT INTO FileEmbeddings (file_id, chunk_index, chunk_type, content_chunk, embedding_vector) VALUES (?, ?, ?, ?, ?)",
                (file_id, -1, 'description', description, json.dumps(desc_vec.tolist()))
            )
            log_execution_step('INGEST_DB', "Saved search priority description embedding.")

            # Handle Expense Table
            if is_expense or category == 'expense':
                log_execution_step('INGEST_DB', f"Registering expense record for file {file_id}")
                cursor.execute("""
                    INSERT INTO Expenses (file_id, amount, currency, category, vendor, date, description)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (
                    file_id,
                    expense_data.get('amount'),
                    expense_data.get('currency', 'USD'),
                    expense_data.get('category', category or 'Uncategorized'),
                    expense_data.get('vendor'),
                    expense_data.get('date'),
                    description
                ))
            
            conn.commit()
            log_execution_step('INGEST_COMPLETE', f"Successfully ingested {final_path.name}")
        except Exception as e:
            log_exception('INGEST_FINAL_DB_ERROR', e, "Faild to complete final metadata updates.")
    
    return {
        "success": True,
        "file_id": file_id,
        "stored_path": str(final_path),
        "description": description,
        "is_expense": is_expense
    }

def search_files(query: str, top_k: int = 5, threshold: float = 0.35) -> List[Dict[str, Any]]:
    """
    Search for files using semantic search (vector embeddings).
    
    The search uses a hybrid approach, boosting matches on the file description
    while also searching within the document content chunks.

    Args:
        query (str): The search query text.
        top_k (int, optional): Number of results to return. Defaults to 5.
        threshold (float, optional): Minimum cosine similarity score to keep a match.
            Valid range is [-1.0, 1.0]. Defaults to 0.35.

    Returns:
        List[Dict]: A list of matching files with 'score', 'snippet', and metadata.
    """
    cda = CommonDataArea()
    try:
        threshold = float(threshold)
    except Exception:
        return [{"error": f"Invalid threshold value: {threshold}"}]
    if threshold < -1.0 or threshold > 1.0:
        return [{"error": f"Threshold out of range [-1.0, 1.0]: {threshold}"}]
    model = get_embedding_model()
    
    if not model:
        return [{"error": "Embedding model not loaded."}]

    try:
        query_vec = model.encode([query])[0]
    except Exception as e:
        return [{"error": f"Encoding failed: {e}"}]

    db_path = _get_embeddings_db_path()
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute("""
        SELECT f.id, f.filename, f.file_path, f.description, 
               fe.chunk_type, fe.content_chunk, fe.embedding_vector 
        FROM FileEmbeddings fe 
        JOIN Files f ON fe.file_id = f.id
    """)
    rows = cursor.fetchall()
    
    try:
        import numpy as np
    except ImportError:
        conn.close()
        return [{"error": "Numpy required for search."}]
    
    results = []
    q_norm = np.linalg.norm(query_vec)
    if q_norm == 0: q_norm = 1e-9

    for row in rows:
        vec_blob = row['embedding_vector']
        if not vec_blob: continue
        
        try:
            vec = np.array(json.loads(vec_blob))
            v_norm = np.linalg.norm(vec)
            if v_norm == 0: v_norm = 1e-9
            
            score = np.dot(query_vec, vec) / (q_norm * v_norm)
            
            # Boost description
            if row['chunk_type'] == 'description':
                score *= 1.2
            
            if score >= threshold:
                results.append({
                    "file_id": row['id'],
                    "filename": row['filename'],
                    "path": row['file_path'],
                    "description": row['description'],
                    "match_type": row['chunk_type'],
                    "snippet": row['content_chunk'][:200],
                    "score": float(score)
                })
        except Exception:
            continue
            
    conn.close()
    results.sort(key=lambda x: x['score'], reverse=True)
    return results[:top_k]

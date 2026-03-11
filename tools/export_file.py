"""Document export tool for CSV/XLSX/DOCX/PDF outputs."""

from __future__ import annotations

import csv
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List
from xml.sax.saxutils import escape

from core.common_data_area import CommonDataArea
from tools.output_utils import (
    build_error_result,
    build_success_result,
    ensure_extension,
    normalize_columns,
    placeholder_text,
    replace_placeholders,
    resolve_output_path,
    rows_as_matrix,
)

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except Exception:
    docx = None
    WD_ORIENT = None
    WD_CELL_VERTICAL_ALIGNMENT = None
    WD_ALIGN_PARAGRAPH = None
    OxmlElement = None
    qn = None
    Inches = None
    Pt = None
    RGBColor = None

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except Exception:
    colors = None
    TA_CENTER = None
    A4 = None
    landscape = None
    ParagraphStyle = None
    getSampleStyleSheet = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None
    Table = None
    TableStyle = None


SUPPORTED_FORMATS = {"csv", "xlsx", "docx", "pdf"}
TEMPLATE_CAPABLE_FORMATS = {"xlsx", "docx"}
_TASK_REPORT_CORE_KEYS = {"id", "title", "status"}
_TASK_REPORT_LAYOUT_WIDTHS = {
    "id": 0.05,
    "title": 0.24,
    "description": 0.30,
    "status": 0.08,
    "scheduled_date": 0.10,
    "due_date": 0.10,
    "priority": 0.06,
    "assigned_to": 0.07,
}
_DOCX_TASK_REPORT_WIDTH_IN = 10.4


def _validate_mode(mode: str) -> str:
    normalized = str(mode or "raw").strip().lower() or "raw"
    if normalized not in {"raw", "template"}:
        raise ValueError("mode must be 'raw' or 'template'.")
    return normalized


def _validate_format(fmt: str) -> str:
    normalized = str(fmt or "").strip().lower()
    if normalized not in SUPPORTED_FORMATS:
        raise ValueError("format must be one of csv, xlsx, docx, pdf.")
    return normalized


def _append_rows_to_docx(document: Any, columns: List[Dict[str, str]], rows: List[Dict[str, Any]] | None) -> None:
    if not columns or not rows:
        return
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, col in enumerate(columns):
        header_cells[idx].text = col["header"]
    for row in rows:
        row_cells = table.add_row().cells
        for idx, col in enumerate(columns):
            row_cells[idx].text = placeholder_text(row.get(col["key"], ""))


def _section_text(section: Dict[str, Any]) -> str:
    return str(section.get("text") or section.get("content") or "")


def _extract_table_section_payload(
    sections: List[Dict[str, Any]] | None,
) -> tuple[List[Dict[str, str]], List[Dict[str, Any]] | None]:
    for section in sections or []:
        if not isinstance(section, dict):
            continue
        if str(section.get("type", "") or "").strip().lower() != "table":
            continue
        section_rows = section.get("rows") if isinstance(section.get("rows"), list) else None
        normalized = normalize_columns(section.get("columns"), section_rows)
        if normalized or section_rows:
            return normalized, section_rows
    return [], None


def _replace_docx_placeholders(document: Any, data: Dict[str, Any] | None) -> int:
    count = 0
    for paragraph in document.paragraphs:
        original = paragraph.text
        updated = replace_placeholders(original, data)
        if updated != original:
            count += 1
            if paragraph.runs:
                paragraph.runs[0].text = updated
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = updated
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                original = cell.text
                updated = replace_placeholders(original, data)
                if updated != original:
                    count += 1
                    cell.text = updated
    return count


def _write_csv(output_path: Path, columns: List[Dict[str, str]], rows: List[Dict[str, Any]] | None) -> Dict[str, Any]:
    if not columns:
        raise ValueError("columns or rows are required for csv export.")
    with open(output_path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([col["header"] for col in columns])
        for row in rows_as_matrix(columns, rows):
            writer.writerow(row)
    return {"row_count": len(rows or []), "column_count": len(columns)}


def _write_xlsx_raw(
    output_path: Path,
    title: str | None,
    sheet_name: str | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
) -> Dict[str, Any]:
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed.")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(sheet_name or "Sheet1")[:31]
    row_offset = 1
    if title:
        ws.cell(row=1, column=1, value=title)
        row_offset = 2
    if columns:
        for idx, col in enumerate(columns, start=1):
            ws.cell(row=row_offset, column=idx, value=col["header"])
        row_offset += 1
    for row_index, row in enumerate(rows or [], start=row_offset):
        for col_index, col in enumerate(columns, start=1):
            ws.cell(row=row_index, column=col_index, value=placeholder_text(row.get(col["key"], "")))
    wb.save(output_path)
    return {"row_count": len(rows or []), "column_count": len(columns), "sheet_names": wb.sheetnames}


def _write_xlsx_template(
    output_path: Path,
    template_path: Path,
    data: Dict[str, Any] | None,
    sheet_name: str | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
) -> Dict[str, Any]:
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed.")
    wb = openpyxl.load_workbook(template_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    placeholder_count = 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                updated = replace_placeholders(cell.value, data)
                if updated != cell.value:
                    placeholder_count += 1
                    cell.value = updated
    if columns and rows:
        start_row = ws.max_row + 2
        for idx, col in enumerate(columns, start=1):
            ws.cell(row=start_row, column=idx, value=col["header"])
        for row_index, row in enumerate(rows or [], start=start_row + 1):
            for col_index, col in enumerate(columns, start=1):
                ws.cell(row=row_index, column=col_index, value=placeholder_text(row.get(col["key"], "")))
    wb.save(output_path)
    return {
        "row_count": len(rows or []),
        "column_count": len(columns),
        "sheet_names": wb.sheetnames,
        "placeholders_filled": placeholder_count,
    }


def _is_task_report_table(columns: List[Dict[str, str]]) -> bool:
    keys = {str(col.get("key", "") or "").strip().lower() for col in columns if isinstance(col, dict)}
    return _TASK_REPORT_CORE_KEYS.issubset(keys) and len(keys.intersection(_TASK_REPORT_LAYOUT_WIDTHS)) >= 5


def _derive_report_title(title: str | None, sections: List[Dict[str, Any]] | None, fallback: str) -> str:
    if str(title or "").strip():
        return str(title).strip()
    for section in sections or []:
        if not isinstance(section, dict):
            continue
        if str(section.get("type", "") or "").strip().lower() == "heading":
            candidate = _section_text(section).strip()
            if candidate:
                return candidate
    return fallback


def _format_task_report_value(key: str, value: Any) -> str:
    text = placeholder_text(value).strip()
    if not text:
        return "-"
    key_name = str(key or "").strip().lower()
    if key_name in {"status", "priority"}:
        return text.title()
    if key_name in {"scheduled_date", "due_date"}:
        return text[:16] if len(text) > 16 else text
    return text


def _task_report_col_widths(columns: List[Dict[str, str]], available_width: float) -> List[float]:
    weights: List[float] = []
    for col in columns:
        key = str(col.get("key", "") or "").strip().lower()
        weights.append(_TASK_REPORT_LAYOUT_WIDTHS.get(key, 0.10))
    total_weight = sum(weights) or 1.0
    return [available_width * (weight / total_weight) for weight in weights]


def _build_task_report_summary(rows: List[Dict[str, Any]] | None) -> str:
    status_counts: Dict[str, int] = {}
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        status = placeholder_text(row.get("status", "Unknown")).strip().title() or "Unknown"
        status_counts[status] = status_counts.get(status, 0) + 1
    parts = [f"Total: {len(rows or [])}"]
    for label in sorted(status_counts):
        parts.append(f"{label}: {status_counts[label]}")
    return " | ".join(parts)


def _docx_clear_paragraph(paragraph: Any) -> Any:
    element = paragraph._element
    for child in list(element):
        element.remove(child)
    return paragraph.add_run()


def _docx_set_cell_text(
    cell: Any,
    text: str,
    *,
    bold: bool = False,
    font_size: int = 9,
    font_color: Any | None = None,
    align: Any | None = None,
) -> None:
    paragraph = cell.paragraphs[0]
    if align is not None and WD_ALIGN_PARAGRAPH is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0) if Pt is not None else None
    paragraph.paragraph_format.space_after = Pt(0) if Pt is not None else None
    paragraph.paragraph_format.keep_together = True
    run = _docx_clear_paragraph(paragraph)
    run.text = text
    run.bold = bold
    if Pt is not None:
        run.font.size = Pt(font_size)
    if font_color is not None and RGBColor is not None:
        run.font.color.rgb = font_color


def _docx_set_cell_shading(cell: Any, fill_hex: str) -> None:
    if OxmlElement is None or qn is None:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill_hex)
    tc_pr.append(shading)


def _docx_set_cell_border(cell: Any, color: str = "C7D2FE", size: str = "6") -> None:
    if OxmlElement is None or qn is None:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _docx_set_repeat_table_header(row: Any) -> None:
    if OxmlElement is None:
        return
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _docx_set_landscape(section: Any) -> None:
    if WD_ORIENT is None or Inches is None:
        return
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)


def _build_task_report_docx(
    document: Any,
    title: str,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
    sections: List[Dict[str, Any]] | None,
    data: Dict[str, Any] | None,
) -> None:
    if docx is None or Inches is None or Pt is None or RGBColor is None:
        raise RuntimeError("python-docx is not installed.")

    _docx_set_landscape(document.sections[0])
    title_rgb = RGBColor(31, 75, 153)
    muted_rgb = RGBColor(75, 85, 99)
    white_rgb = RGBColor(255, 255, 255)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = title_rgb

    summary_paragraph = document.add_paragraph()
    summary_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary_run = summary_paragraph.add_run(_build_task_report_summary(rows))
    summary_run.font.size = Pt(9)
    summary_run.font.color.rgb = muted_rgb

    generated_paragraph = document.add_paragraph()
    generated_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated_run = generated_paragraph.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    generated_run.font.size = Pt(9)
    generated_run.font.color.rgb = muted_rgb

    consumed_heading = False
    for section in sections or []:
        if not isinstance(section, dict):
            continue
        section_type = str(section.get("type", "") or "").strip().lower()
        section_text = replace_placeholders(_section_text(section), data).strip()
        if section_type == "heading" and not consumed_heading and section_text == title:
            consumed_heading = True
            continue
        if section_type == "paragraph" and section_text:
            paragraph = document.add_paragraph(section_text)
            paragraph.paragraph_format.space_after = Pt(6)
        elif section_type == "heading" and section_text:
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(4)
            heading.paragraph_format.space_after = Pt(3)
            run = heading.add_run(section_text)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = title_rgb

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    widths = _task_report_col_widths(columns, _DOCX_TASK_REPORT_WIDTH_IN)
    header_row = table.rows[0]
    _docx_set_repeat_table_header(header_row)
    for idx, col in enumerate(columns):
        cell = header_row.cells[idx]
        cell.width = Inches(widths[idx])
        if WD_CELL_VERTICAL_ALIGNMENT is not None:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _docx_set_cell_shading(cell, "1F4B99")
        _docx_set_cell_border(cell, color="1E3A8A", size="10")
        _docx_set_cell_text(
            cell,
            col["header"],
            bold=True,
            font_size=8,
            font_color=white_rgb,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_index, row in enumerate(rows or [], start=1):
        if not isinstance(row, dict):
            continue
        table_row = table.add_row()
        fill = "F8FAFC" if row_index % 2 == 0 else "FFFFFF"
        for idx, col in enumerate(columns):
            cell = table_row.cells[idx]
            cell.width = Inches(widths[idx])
            if WD_CELL_VERTICAL_ALIGNMENT is not None:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _docx_set_cell_shading(cell, fill)
            _docx_set_cell_border(cell, color="C7D2FE", size="6")
            align = WD_ALIGN_PARAGRAPH.CENTER if col["key"] == "id" else WD_ALIGN_PARAGRAPH.LEFT
            _docx_set_cell_text(
                cell,
                _format_task_report_value(col["key"], row.get(col["key"], "")),
                font_size=8,
                align=align,
            )


def _write_docx_raw(
    output_path: Path,
    title: str | None,
    data: Dict[str, Any] | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
    sections: List[Dict[str, Any]] | None,
) -> Dict[str, Any]:
    if docx is None:
        raise RuntimeError("python-docx is not installed.")
    document = docx.Document()

    effective_columns = columns
    effective_rows = rows
    if (not effective_columns and not effective_rows) and sections:
        effective_columns, effective_rows = _extract_table_section_payload(sections)

    layout = "generic"
    orientation = "portrait"
    report_title = _derive_report_title(title, sections, "Task Report")

    if effective_columns and _is_task_report_table(effective_columns):
        _build_task_report_docx(document, report_title, effective_columns, effective_rows, sections, data)
        layout = "task_report"
        orientation = "landscape"
    else:
        if title:
            document.add_heading(title, level=1)
        for section in sections or []:
            if not isinstance(section, dict):
                continue
            section_type = str(section.get("type", "paragraph") or "paragraph").strip().lower()
            section_text = replace_placeholders(_section_text(section), data)
            if section_type == "paragraph":
                if section_text:
                    document.add_paragraph(section_text)
            elif section_type == "heading":
                if section_text:
                    document.add_heading(section_text, level=int(section.get("level", 2) or 2))
            elif section_type == "table":
                table_rows = section.get("rows") if isinstance(section.get("rows"), list) else effective_rows
                table_columns = normalize_columns(section.get("columns"), table_rows) or effective_columns
                _append_rows_to_docx(document, table_columns, table_rows)
        if data and not sections:
            for key, value in data.items():
                document.add_paragraph(f"{key}: {placeholder_text(value)}")
        if effective_columns and effective_rows and not any(
            str((section or {}).get("type", "")).strip().lower() == "table" for section in sections or []
        ):
            _append_rows_to_docx(document, effective_columns, effective_rows)

    document.save(output_path)
    return {
        "row_count": len(effective_rows or []),
        "column_count": len(effective_columns),
        "placeholders_filled": 0,
        "layout": layout,
        "page_orientation": orientation,
    }


def _write_docx_template(
    output_path: Path,
    template_path: Path,
    data: Dict[str, Any] | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
) -> Dict[str, Any]:
    if docx is None:
        raise RuntimeError("python-docx is not installed.")
    document = docx.Document(str(template_path))
    placeholder_count = _replace_docx_placeholders(document, data)
    _append_rows_to_docx(document, columns, rows)
    document.save(output_path)
    return {"row_count": len(rows or []), "column_count": len(columns), "placeholders_filled": placeholder_count}


def _pdf_text(value: Any) -> str:
    return escape(placeholder_text(value)).replace("\n", "<br/>")


def _build_generic_pdf_story(
    title: str | None,
    data: Dict[str, Any] | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
    sections: List[Dict[str, Any]] | None,
) -> List[Any]:
    if getSampleStyleSheet is None or Paragraph is None or Spacer is None:
        raise RuntimeError("reportlab is not installed.")
    styles = getSampleStyleSheet()
    story: List[Any] = []
    if title:
        story.append(Paragraph(title, styles["Title"]))
        story.append(Spacer(1, 12))

    effective_columns = columns
    effective_rows = rows
    if (not effective_columns and not effective_rows) and sections:
        effective_columns, effective_rows = _extract_table_section_payload(sections)

    table_rendered = False
    for section in sections or []:
        if not isinstance(section, dict):
            continue
        section_type = str(section.get("type", "paragraph") or "paragraph").strip().lower()
        section_text = replace_placeholders(_section_text(section), data)
        if section_type == "heading":
            if section_text:
                story.append(Paragraph(section_text, styles["Heading2"]))
                story.append(Spacer(1, 8))
        elif section_type == "paragraph":
            if section_text:
                story.append(Paragraph(section_text, styles["BodyText"]))
                story.append(Spacer(1, 8))
        elif section_type == "table" and Table is not None and TableStyle is not None and colors is not None:
            table_rows = section.get("rows") if isinstance(section.get("rows"), list) else effective_rows
            table_columns = normalize_columns(section.get("columns"), table_rows) or effective_columns
            if table_columns:
                table_data = [[col["header"] for col in table_columns]] + rows_as_matrix(table_columns, table_rows)
                table = Table(table_data, repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9e8fb")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                story.append(table)
                story.append(Spacer(1, 8))
                table_rendered = True
    if data and not sections:
        for key, value in data.items():
            story.append(Paragraph(f"<b>{key}</b>: {placeholder_text(value)}", styles["BodyText"]))
            story.append(Spacer(1, 4))
    if effective_columns and not table_rendered and Table is not None and TableStyle is not None and colors is not None:
        table_data = [[col["header"] for col in effective_columns]] + rows_as_matrix(effective_columns, effective_rows)
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9e8fb")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)
    return story


def _build_task_report_pdf_story(
    title: str | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
    page_width: float,
) -> List[Any]:
    if getSampleStyleSheet is None or Paragraph is None or Spacer is None or ParagraphStyle is None:
        raise RuntimeError("reportlab is not installed.")

    styles = getSampleStyleSheet()
    story: List[Any] = []
    report_title = title or "Task Report"
    meta_style = ParagraphStyle(
        "TaskReportMeta",
        parent=styles["BodyText"],
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#4b5563"),
        spaceAfter=6,
    )
    header_style = ParagraphStyle(
        "TaskReportHeader",
        parent=styles["BodyText"],
        fontSize=8,
        leading=10,
        alignment=TA_CENTER,
        textColor=colors.white,
    )
    cell_style = ParagraphStyle(
        "TaskReportCell",
        parent=styles["BodyText"],
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#111827"),
    )

    story.append(Paragraph(_pdf_text(report_title), styles["Title"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(_pdf_text(_build_task_report_summary(rows)), meta_style))
    story.append(Paragraph(_pdf_text(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"), meta_style))
    story.append(Spacer(1, 8))

    table_rows: List[List[Any]] = []
    table_rows.append([Paragraph(_pdf_text(col["header"]), header_style) for col in columns])
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        rendered_row = []
        for col in columns:
            key = col["key"]
            rendered_row.append(Paragraph(_pdf_text(_format_task_report_value(key, row.get(key, ""))), cell_style))
        table_rows.append(rendered_row)

    col_widths = _task_report_col_widths(columns, page_width)
    table = Table(table_rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4b99")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c7d2fe")),
                ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.HexColor("#1e3a8a")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]
        )
    )
    story.append(table)
    return story


def _write_pdf(
    output_path: Path,
    title: str | None,
    data: Dict[str, Any] | None,
    columns: List[Dict[str, str]],
    rows: List[Dict[str, Any]] | None,
    sections: List[Dict[str, Any]] | None,
) -> Dict[str, Any]:
    if SimpleDocTemplate is None or A4 is None or landscape is None:
        raise RuntimeError("reportlab is not installed.")

    effective_columns = columns
    effective_rows = rows
    if (not effective_columns and not effective_rows) and sections:
        effective_columns, effective_rows = _extract_table_section_payload(sections)

    layout = "generic"
    page_size = A4
    left_margin = right_margin = 36
    top_margin = 36
    bottom_margin = 28

    if effective_columns and _is_task_report_table(effective_columns):
        layout = "task_report"
        page_size = landscape(A4)
        left_margin = right_margin = 18
        top_margin = 24
        bottom_margin = 18

    available_width = page_size[0] - left_margin - right_margin
    if layout == "task_report":
        story = _build_task_report_pdf_story(title, effective_columns, effective_rows, available_width)
        orientation = "landscape"
    else:
        story = _build_generic_pdf_story(title, data, effective_columns, effective_rows, sections)
        orientation = "portrait"

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
    )
    doc.build(story)
    output_path.write_bytes(buffer.getvalue())
    return {
        "row_count": len(effective_rows or []),
        "column_count": len(effective_columns),
        "placeholders_filled": 0,
        "layout": layout,
        "page_orientation": orientation,
    }


def export_file(
    output_filename: str,
    format: str,
    mode: str = "raw",
    title: str | None = None,
    template_path: str | None = None,
    data: Dict[str, Any] | None = None,
    columns: List[Dict[str, Any]] | None = None,
    rows: List[Dict[str, Any]] | None = None,
    sections: List[Dict[str, Any]] | None = None,
    sheet_name: str | None = None,
    subfolder: str = "generated",
    options: Dict[str, Any] | None = None,
    status_callback=None,
) -> Dict[str, Any]:
    """Create a document file under the configured file storage path."""

    cda = CommonDataArea()
    output_path: Path | None = None
    warnings: List[str] = []
    try:
        fmt = _validate_format(format)
        run_mode = _validate_mode(mode)
        _, output_path = resolve_output_path(output_filename, subfolder, cda)
        output_path = ensure_extension(output_path, fmt)
        normalized_columns = normalize_columns(columns, rows)
        meta: Dict[str, Any] = {
            "title": str(title or ""),
            "sheet_names": [],
            "row_count": len(rows or []),
            "placeholder_count": len(data or {}),
            "placeholders_filled": 0,
            "template_used": "",
            "storage_subfolder": subfolder,
        }

        if run_mode == "template":
            if fmt not in TEMPLATE_CAPABLE_FORMATS:
                raise ValueError(f"Template mode is not supported for {fmt}.")
            if not template_path:
                raise ValueError("template_path is required when mode='template'.")
            template = Path(str(template_path)).resolve()
            if not template.exists():
                raise FileNotFoundError("Template file does not exist.")
            meta["template_used"] = str(template)
            if status_callback:
                status_callback(f"Filling {fmt.upper()} template...")
            if fmt == "xlsx":
                details = _write_xlsx_template(output_path, template, data, sheet_name, normalized_columns, rows)
            else:
                details = _write_docx_template(output_path, template, data, normalized_columns, rows)
        else:
            if status_callback:
                status_callback(f"Creating {fmt.upper()} file...")
            if fmt == "csv":
                details = _write_csv(output_path, normalized_columns, rows)
            elif fmt == "xlsx":
                details = _write_xlsx_raw(output_path, title, sheet_name, normalized_columns, rows)
            elif fmt == "docx":
                details = _write_docx_raw(output_path, title, data, normalized_columns, rows, sections)
            else:
                details = _write_pdf(output_path, title, data, normalized_columns, rows, sections)

        meta.update(details)
        if options:
            meta["options"] = options
        if status_callback:
            status_callback("")
        return build_success_result(
            output_path=output_path,
            fmt=fmt,
            mode=run_mode,
            message="Document exported successfully.",
            warnings=warnings,
            meta=meta,
        )
    except Exception as exc:
        if status_callback:
            status_callback("")
        return build_error_result(
            output_path=output_path,
            fmt=str(format or ""),
            mode=str(mode or "raw"),
            message="Export failed.",
            error=str(exc),
            warnings=warnings,
            meta={"storage_subfolder": subfolder},
        )

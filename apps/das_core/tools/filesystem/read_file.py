"""Read file tool."""

from __future__ import annotations

import os
from typing import Dict

try:
    import docx
except Exception:  # noqa: BLE001
    docx = None

try:
    import openpyxl
except Exception:  # noqa: BLE001
    openpyxl = None

try:
    import pptx
except Exception:  # noqa: BLE001
    pptx = None

from PyPDF2 import PdfReader

from tools.filesystem import ensure_accessible


def _read_docx(file_path: str) -> str:
    if docx is None:
        raise RuntimeError('python-docx is not installed.')
    document = docx.Document(file_path)
    return '\n'.join(para.text for para in document.paragraphs)


def _read_xlsx(file_path: str) -> str:
    if openpyxl is None:
        raise RuntimeError('openpyxl is not installed.')
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    parts: list[str] = []
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        parts.append(f'Sheet: {sheet_name}')
        for row in worksheet.iter_rows(values_only=True):
            values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if values:
                parts.append(' | '.join(values))
    return '\n'.join(parts)


def _read_pptx(file_path: str) -> str:
    if pptx is None:
        raise RuntimeError('python-pptx is not installed.')
    presentation = pptx.Presentation(file_path)
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            text = getattr(shape, 'text', '')
            if text:
                parts.append(text)
    return '\n'.join(parts)


def read_file(file_path: str) -> Dict:
    """
    Read the content of a file (text, PDF, DOCX, XLSX, or PPTX).

    Args:
        file_path (str): The absolute path to the file to read.

    Returns:
        Dict: A dictionary containing 'success' (bool) and 'data' (dict) or 'error' (str).
              Data includes 'content' (str) and 'type' (str).
    """

    ok, error = ensure_accessible(file_path)
    if not ok:
        return {'success': False, 'error': error}

    if not os.path.exists(file_path):
        return {'success': False, 'error': 'File does not exist.'}

    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == '.pdf':
            reader = PdfReader(file_path)
            text_parts = []
            for page in reader.pages:
                text_parts.append(page.extract_text() or '')
            content = '\n'.join(text_parts)
            return {'success': True, 'data': {'file_path': file_path, 'type': 'pdf', 'content': content}}

        if ext in {'.doc', '.docx'}:
            content = _read_docx(file_path)
            return {'success': True, 'data': {'file_path': file_path, 'type': 'docx', 'content': content}}

        if ext in {'.xls', '.xlsx'}:
            content = _read_xlsx(file_path)
            return {'success': True, 'data': {'file_path': file_path, 'type': 'xlsx', 'content': content}}

        if ext in {'.ppt', '.pptx'}:
            content = _read_pptx(file_path)
            return {'success': True, 'data': {'file_path': file_path, 'type': 'pptx', 'content': content}}

        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        return {'success': True, 'data': {'file_path': file_path, 'type': 'text', 'content': content}}
    except Exception as exc:  # noqa: BLE001
        return {'success': False, 'error': str(exc)}
